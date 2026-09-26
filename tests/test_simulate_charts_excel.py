import openpyxl
import pytest

from corefin.assumptions.schema import RootConfig
from corefin.optimize.structure import build_structure
from corefin.simulate.attribution import (
    compute_value_creation_bridge,
    summarize_value_creation_bridge,
)
from corefin.simulate.charts import (
    render_breach_distress_chart,
    render_irr_histogram,
    render_leverage_fan_chart,
    render_structure_comparison_chart,
    render_tornado_chart,
    render_value_bridge_waterfall,
)
from corefin.simulate.compare import compare_structures, input_config_decision_values
from corefin.simulate.excel_export import export_simulation_to_excel
from corefin.simulate.importance import compute_driver_importance, compute_tornado_chart
from corefin.timeline import Timeline
from tests.conftest import minimal_config_dict


def _config() -> RootConfig:
    data = minimal_config_dict(n_periods=6)
    data["covenants"] = [
        {"name": "Max Lev", "metric": "total_net_leverage", "threshold": 4.0, "test_from_period": 0}
    ]
    data["optimizer"] = {
        "decision_variables": [{"tranche_name": "TLB", "min_multiple": 1.0, "max_multiple": 5.0}],
    }
    data["scenario"]["driver_vol"] = {
        "revenue_growth_std": 0.06,
        "ebitda_margin_std": 0.03,
        "exit_multiple_std": 0.8,
        "base_rate_std": 0.01,
    }
    data["simulate"] = {
        "stress_scenarios": [
            {"name": "Recession", "shocks": {"recession": {"start_year_index": 2}}},
        ]
    }
    return RootConfig.model_validate(data)


@pytest.fixture(scope="module")
def comparison_bundle():
    config = _config()
    timeline = Timeline.annual(n_periods=6)
    named = {
        "Input": input_config_decision_values(config),
        "Optimized": build_structure(config, {"TLB": 2.5}).decision_values,
    }
    entries = compare_structures(config, timeline, named, n_scenarios=1500, seed=5)
    primary = entries[-1]
    bridge = compute_value_creation_bridge(config, primary.simulation)
    bridge_summaries = summarize_value_creation_bridge(bridge, primary.simulation.exit_result.irr)
    importance = compute_driver_importance(primary.simulation)
    tornado = compute_tornado_chart(config, primary.candidate, timeline, primary.simulation)
    return config, timeline, entries, primary, bridge_summaries, importance, tornado


def test_render_irr_histogram_writes_png(tmp_path, comparison_bundle):
    _, _, _, primary, _, _, _ = comparison_bundle
    path = tmp_path / "irr_histogram.png"
    render_irr_histogram(primary.downside, primary.simulation.exit_result.irr, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_render_leverage_fan_chart_writes_png(tmp_path, comparison_bundle):
    _, timeline, _, primary, _, _, _ = comparison_bundle
    path = tmp_path / "leverage_fan.png"
    render_leverage_fan_chart(primary.downside, timeline, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_render_breach_distress_chart_writes_png(tmp_path, comparison_bundle):
    _, timeline, _, primary, _, _, _ = comparison_bundle
    path = tmp_path / "breach_distress.png"
    render_breach_distress_chart(primary.downside, timeline, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_render_value_bridge_waterfall_writes_png(tmp_path, comparison_bundle):
    _, _, _, _, bridge_summaries, _, _ = comparison_bundle
    path = tmp_path / "value_bridge.png"
    render_value_bridge_waterfall(bridge_summaries[0], str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_render_tornado_chart_writes_png(tmp_path, comparison_bundle):
    _, _, _, _, _, _, tornado = comparison_bundle
    path = tmp_path / "tornado.png"
    render_tornado_chart(tornado, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_render_structure_comparison_chart_writes_png(tmp_path, comparison_bundle):
    _, _, entries, _, _, _, _ = comparison_bundle
    path = tmp_path / "structure_comparison.png"
    render_structure_comparison_chart(entries, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_export_simulation_to_excel_writes_all_expected_sheets(tmp_path, comparison_bundle):
    config, timeline, entries, primary, bridge_summaries, importance, tornado = comparison_bundle
    path = tmp_path / "simulation.xlsx"
    export_simulation_to_excel(
        str(path), timeline, entries, primary.name, bridge_summaries, importance, tornado
    )
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    expected_sheets = {
        "Summary",
        "Distributions",
        "Stress Scenarios",
        "Breach & Distress by Year",
        "Attribution",
        "Driver Importance",
        "Structure Comparison",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_summary_sheet_has_one_column_per_structure(tmp_path, comparison_bundle):
    config, timeline, entries, primary, bridge_summaries, importance, tornado = comparison_bundle
    path = tmp_path / "simulation.xlsx"
    export_simulation_to_excel(
        str(path), timeline, entries, primary.name, bridge_summaries, importance, tornado
    )
    wb = openpyxl.load_workbook(path)
    ws = wb["Summary"]
    header = [c.value for c in ws[1]]
    assert set(header[1:]) == {e.name for e in entries}


def test_attribution_sheet_has_four_bridge_rows(tmp_path, comparison_bundle):
    config, timeline, entries, primary, bridge_summaries, importance, tornado = comparison_bundle
    path = tmp_path / "simulation.xlsx"
    export_simulation_to_excel(
        str(path), timeline, entries, primary.name, bridge_summaries, importance, tornado
    )
    wb = openpyxl.load_workbook(path)
    ws = wb["Attribution"]
    header = [c.value for c in ws[1]]
    assert set(header[1:]) == {"Average", "P10", "Median", "P90"}
