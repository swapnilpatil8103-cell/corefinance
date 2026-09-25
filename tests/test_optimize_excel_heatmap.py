import openpyxl
import pytest

from corefin.assumptions.schema import RootConfig
from corefin.metrics.credit_metrics import compute_credit_metrics
from corefin.optimize.diagnostics import compute_diagnostics
from corefin.optimize.excel_export import export_optimization_to_excel
from corefin.optimize.heatmap import render_leverage_frontier, render_objective_heatmap
from corefin.optimize.search import run_optimization
from corefin.timeline import Timeline
from tests.test_debt_integration import debt_config_dict


def _optimizer_config(grid_points: int = 4) -> RootConfig:
    data = debt_config_dict(n_periods=5)
    data["optimizer"] = {
        "decision_variables": [
            {"tranche_name": "TLB", "min_multiple": 0.5, "max_multiple": 4.0},
            {"tranche_name": "Notes", "min_multiple": 0.0, "max_multiple": 1.5},
        ],
        "search": {
            "grid_points_per_dimension": grid_points,
            "n_scenarios_search": 300,
            "n_scenarios_confirm": 600,
            "refine_grid_points_per_dimension": 3,
        },
    }
    return RootConfig.model_validate(data)


@pytest.fixture(scope="module")
def optimization_result():
    config = _optimizer_config()
    timeline = Timeline.annual(n_periods=5)
    result = run_optimization(config, timeline)
    diagnostics = compute_diagnostics(config, timeline, result)
    return config, timeline, result, diagnostics


def _recommended_credit_metrics(result):
    confirmation = result.confirmation
    return compute_credit_metrics(
        confirmation.result.debt_schedule,
        confirmation.candidate.tranches,
        confirmation.result.income_statement.ebitda,
        confirmation.result.cash_flow_statement.capex,
        confirmation.result.balance_sheet.cash,
        confirmation.result.balance_sheet.total_debt,
    )


def test_export_optimization_to_excel_writes_all_expected_sheets(tmp_path, optimization_result):
    config, timeline, result, diagnostics = optimization_result
    credit_metrics = _recommended_credit_metrics(result)
    path = tmp_path / "optimization.xlsx"
    export_optimization_to_excel(str(path), timeline, result, diagnostics, credit_metrics)

    assert path.exists()
    wb = openpyxl.load_workbook(path)
    expected_sheets = {
        "Grid Results",
        "Efficient Frontier",
        "Income Statement",
        "Balance Sheet",
        "Cash Flow Statement",
        "Debt Schedule",
        "Credit Metrics",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_grid_results_sheet_has_one_row_per_evaluation(tmp_path, optimization_result):
    config, timeline, result, diagnostics = optimization_result
    credit_metrics = _recommended_credit_metrics(result)
    path = tmp_path / "optimization.xlsx"
    export_optimization_to_excel(str(path), timeline, result, diagnostics, credit_metrics)

    wb = openpyxl.load_workbook(path)
    ws = wb["Grid Results"]
    n_grid = result.grid_result.n_evaluated
    n_refine = result.refinement_result.n_evaluated if result.refinement_result else 0
    assert ws.max_row == 1 + n_grid + n_refine  # header + every evaluation


def test_efficient_frontier_only_contains_feasible_and_marks_recommended(
    tmp_path, optimization_result
):
    config, timeline, result, diagnostics = optimization_result
    credit_metrics = _recommended_credit_metrics(result)
    path = tmp_path / "optimization.xlsx"
    export_optimization_to_excel(str(path), timeline, result, diagnostics, credit_metrics)

    wb = openpyxl.load_workbook(path)
    ws = wb["Efficient Frontier"]
    header = [c.value for c in ws[1]]
    feasible_col = header.index("Feasible") + 1
    recommended_col = header.index("Recommended") + 1
    values = [row[feasible_col - 1].value for row in ws.iter_rows(min_row=2)]
    assert all(v for v in values)  # every row is feasible
    recommended_flags = [row[recommended_col - 1].value for row in ws.iter_rows(min_row=2)]
    assert sum(1 for v in recommended_flags if v) >= 1


def test_render_objective_heatmap_writes_png(tmp_path, optimization_result):
    config, timeline, result, diagnostics = optimization_result
    path = tmp_path / "heatmap.png"
    render_objective_heatmap(result, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000  # a real image, not an empty/broken file


def test_render_objective_heatmap_requires_two_decision_variables(tmp_path):
    data = debt_config_dict(n_periods=5)
    data["optimizer"] = {
        "decision_variables": [{"tranche_name": "TLB", "min_multiple": 1.0, "max_multiple": 3.0}],
        "search": {"grid_points_per_dimension": 3, "n_scenarios_search": 100},
    }
    config = RootConfig.model_validate(data)
    timeline = Timeline.annual(n_periods=5)
    result = run_optimization(config, timeline)
    with pytest.raises(ValueError, match="two decision variables"):
        render_objective_heatmap(result, str(tmp_path / "heatmap.png"))


def test_render_leverage_frontier_writes_png(tmp_path, optimization_result):
    config, timeline, result, diagnostics = optimization_result
    path = tmp_path / "leverage_frontier.png"
    render_leverage_frontier(result, str(path))
    assert path.exists()
    assert path.stat().st_size > 1000


def test_binding_and_relaxation_sheets_present(tmp_path, optimization_result):
    config, timeline, result, diagnostics = optimization_result
    credit_metrics = _recommended_credit_metrics(result)
    path = tmp_path / "optimization.xlsx"
    export_optimization_to_excel(str(path), timeline, result, diagnostics, credit_metrics)

    wb = openpyxl.load_workbook(path)
    assert "Binding & Limiting" in wb.sheetnames
    assert "Relaxation Sensitivity" in wb.sheetnames
